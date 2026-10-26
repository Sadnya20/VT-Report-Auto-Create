
import streamlit as st
import pandas as pd
import io
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

st.set_page_config(
    page_title="Employee Attendance Report",
    page_icon="📊",
    layout="wide"
)

REQUIRED_COLUMNS = [
    "Date",
    "Employee Name",
    "Designation",
    "Working_With",
    "Covered_Beat",
    "Attendance",
    "Market Objective",
    "Attendance Time",
    "Atten. Remarks"
]

ALLOWED_DESIGNATIONS = [
    "ASM", "ASE", "SDM", "RSM",
    "TSE", "TSA", "NSM", "ZSM"
]

st.title("📊 Employee Attendance Report")
st.write("Upload Excel or CSV to generate your report.")

uploaded_file = st.file_uploader(
    "Upload Excel or CSV",
    type=["xlsx", "xls", "csv"]
)

if uploaded_file is not None:
    try:
        filename = uploaded_file.name.lower()

        # Read source file
        if filename.endswith(".csv"):
            df = pd.read_csv(uploaded_file)
        else:
            excel_file = pd.ExcelFile(uploaded_file)

            if len(excel_file.sheet_names) > 1:
                sheet = st.selectbox(
                    "Select worksheet",
                    excel_file.sheet_names
                )
            else:
                sheet = excel_file.sheet_names[0]

            df = pd.read_excel(
                uploaded_file,
                sheet_name=sheet
            )

        # Clean column header spaces
        df.columns = df.columns.astype(str).str.strip()

        st.success("File uploaded successfully!")

        # Check required columns
        missing_columns = [
            col for col in REQUIRED_COLUMNS
            if col not in df.columns
        ]

        if missing_columns:
            st.error(
                "Missing columns: "
                + ", ".join(missing_columns)
            )
            st.stop()

        # Keep only required columns
        report = df[REQUIRED_COLUMNS].copy()

        # REMOVE rows where Atten. Remarks is None,
        # blank, missing, or text "None"
        remarks = (
            report["Atten. Remarks"]
            .astype("string")
            .str.strip()
            .str.casefold()
        )

        report = report[
            remarks.notna()
            & ~remarks.isin(["none", ""])
        ].copy()

        # Standardize Designation
        report["Designation"] = (
            report["Designation"]
            .astype("string")
            .str.strip()
            .str.upper()
        )

        # Keep only allowed designations
        report = report[
            report["Designation"].isin(
                ALLOWED_DESIGNATIONS
            )
        ].copy()

        st.subheader("Filters")

        c1, c2 = st.columns(2)

        with c1:
            present_designations = [
                d for d in ALLOWED_DESIGNATIONS
                if d in report["Designation"].values
            ]

            designations = st.multiselect(
                "Select Designation",
                present_designations,
                default=present_designations
            )

        with c2:
            employees = sorted(
                report["Employee Name"]
                .dropna()
                .astype(str)
                .unique()
                .tolist()
            )

            selected_employees = st.multiselect(
                "Select Employee",
                employees
            )

        # Apply filters
        report = report[
            report["Designation"].isin(designations)
        ].copy()

        if selected_employees:
            report = report[
                report["Employee Name"]
                .astype(str)
                .isin(selected_employees)
            ].copy()

        st.subheader("Report Summary")

        c1, c2, c3 = st.columns(3)

        c1.metric("Original Rows", len(df))
        c2.metric("Final Report Rows", len(report))
        c3.metric(
            "Unique Employees",
            report["Employee Name"].nunique()
        )

        st.subheader("Report Preview")

        st.dataframe(
            report,
            use_container_width=True,
            hide_index=True
        )

        # Designation-wise summary
        summary = (
            report.groupby("Designation")
            .size()
            .reindex(ALLOWED_DESIGNATIONS, fill_value=0)
            .rename_axis("Designation")
            .reset_index(name="Record Count")
        )

        # Generate Excel report
        output = io.BytesIO()

        with pd.ExcelWriter(
            output,
            engine="openpyxl"
        ) as writer:

            report.to_excel(
                writer,
                sheet_name="Attendance Report",
                index=False
            )

            summary.to_excel(
                writer,
                sheet_name="Designation Summary",
                index=False
            )

            # Format both sheets
            for ws in writer.book.worksheets:
                ws.freeze_panes = "A2"
                ws.sheet_view.showGridLines = False

                if ws.max_row >= 1 and ws.max_column >= 1:
                    ws.auto_filter.ref = ws.dimensions

                for cell in ws[1]:
                    cell.font = Font(
                        bold=True,
                        color="FFFFFF"
                    )
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="17365D"
                    )
                    cell.alignment = Alignment(
                        vertical="center",
                        wrap_text=True
                    )

                ws.row_dimensions[1].height = 30

                for column_cells in ws.columns:
                    max_length = max(
                        (
                            len(str(cell.value))
                            for cell in list(column_cells)[:500]
                            if cell.value is not None
                        ),
                        default=10
                    )

                    col_letter = get_column_letter(
                        column_cells[0].column
                    )

                    ws.column_dimensions[col_letter].width = min(
                        max(max_length + 2, 12),
                        35
                    )

        output.seek(0)

        st.download_button(
            label="📥 Download Formatted Excel Report",
            data=output.getvalue(),
            file_name="Employee_Attendance_Report.xlsx",
            mime=(
                "application/vnd.openxmlformats-officedocument"
                ".spreadsheetml.sheet"
            ),
            type="primary"
        )

    except Exception as e:
        st.error(f"Error processing file: {e}")
